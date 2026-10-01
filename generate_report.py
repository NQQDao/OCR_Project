# -*- coding: utf-8 -*-
"""
generate_report.py
Tạo báo cáo tổng quan dự án OCR Nhật ký xẻ gỗ dưới dạng DOCX.
Chạy: python generate_report.py
Output: bao_cao_tong_quan_ocr_go.docx
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

# ─────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────

def set_cell_bg(cell, hex_color: str):
    """Đặt màu nền cho ô trong bảng."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def cell_text(cell, text, bold=False, size=10, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    """Ghi text vào ô, áp dụng định dạng."""
    para = cell.paragraphs[0]
    para.alignment = align
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return h


def add_paragraph(doc, text, bold=False, size=11, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.7)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return p


def add_bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.size = Pt(size)
    return p


def add_table_with_header(doc, headers, rows, col_widths=None, header_color="1F4E79"):
    """Tạo bảng có header màu đậm."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        set_cell_bg(cell, header_color)
        cell_text(cell, h, bold=True, size=10, color=(255, 255, 255),
                  align=WD_ALIGN_PARAGRAPH.CENTER)

    # Data rows
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            if r_idx % 2 == 1:
                set_cell_bg(cell, "DEEAF1")
            cell_text(cell, str(val), size=10)

    # Column widths
    if col_widths:
        for r in table.rows:
            for i, cell in enumerate(r.cells):
                cell.width = Cm(col_widths[i])

    return table


# ─────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────

def build_report():
    doc = Document()

    # ── Page margin ──
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)

    # ── Default font ──
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    # ═══════════════════════════════════════════
    # TRANG BÌA
    # ═══════════════════════════════════════════
    doc.add_paragraph()
    doc.add_paragraph()

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_r = title_p.add_run("BÁO CÁO TỔNG QUAN")
    title_r.bold = True
    title_r.font.size = Pt(22)
    title_r.font.color.rgb = RGBColor(31, 78, 121)
    title_r.font.name = "Times New Roman"

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_r = sub_p.add_run("DỰ ÁN HỆ THỐNG OCR NHẬT KÝ XẺ GỖ")
    sub_r.bold = True
    sub_r.font.size = Pt(16)
    sub_r.font.color.rgb = RGBColor(68, 114, 196)
    sub_r.font.name = "Times New Roman"

    doc.add_paragraph()
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_r = date_p.add_run(f"Ngày lập: {datetime.date.today().strftime('%d/%m/%Y')}")
    date_r.font.size = Pt(12)
    date_r.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_page_break()

    # ═══════════════════════════════════════════
    # 1. TỔNG QUAN DỰ ÁN
    # ═══════════════════════════════════════════
    add_heading(doc, "1. TỔNG QUAN DỰ ÁN", level=1)

    add_paragraph(doc,
        "Hệ thống OCR Nhật ký xẻ gỗ là giải pháp số hoá tự động các phiếu nhật ký "
        "xẻ gỗ bằng công nghệ AI Vision (OCR). Người dùng chụp ảnh phiếu thủ công, "
        "upload lên hệ thống, AI sẽ nhận dạng và trích xuất dữ liệu vào cơ sở dữ liệu "
        "có cấu trúc, hỗ trợ kiểm tra đối chiếu và xuất báo cáo Excel theo lô.")

    add_paragraph(doc, "Mục tiêu:", bold=True)
    for item in [
        "Số hoá 100% phiếu nhật ký xẻ gỗ, loại bỏ nhập liệu thủ công.",
        "Đảm bảo tính chính xác dữ liệu qua bước đối chiếu trước khi xuất Excel.",
        "Quản lý đơn giá gỗ theo mã gỗ tự động từ cơ sở dữ liệu.",
        "Hỗ trợ phân tích tỷ lệ thu hồi (m³ thành khí / m³ gỗ tròn).",
    ]:
        add_bullet(doc, item)

    doc.add_paragraph()

    # ═══════════════════════════════════════════
    # 2. KIẾN TRÚC HỆ THỐNG
    # ═══════════════════════════════════════════
    add_heading(doc, "2. KIẾN TRÚC HỆ THỐNG", level=1)
    add_paragraph(doc,
        "Dự án được triển khai hoàn toàn Serverless trên nền tảng Cloudflare, "
        "không cần máy chủ riêng, chi phí vận hành thấp và khả năng mở rộng cao.")

    arch_rows = [
        ("Cloudflare Workers", "Runtime serverless – xử lý toàn bộ API backend (OCR, DB, Excel)"),
        ("Cloudflare D1 (SQLite)", "Cơ sở dữ liệu quan hệ – lưu phiếu, chi tiết gỗ, đơn giá"),
        ("Cloudflare R2", "Object Storage – lưu trữ ảnh phiếu gốc upload"),
        ("OpenRouter API", "Multi-model AI Gateway – ưu tiên Gemini 2.5 Flash, fallback tự động"),
        ("Google Gemini Vision", "Fallback OCR khi OpenRouter không khả dụng"),
        ("ExcelJS (Node.js)", "Tạo file Excel báo cáo nhiều sheet từ dữ liệu D1"),
        ("HTML/JS Frontend", "Giao diện SPA nhúng trong Worker – không cần server riêng"),
    ]
    add_table_with_header(doc,
        ["Thành phần", "Vai trò"],
        arch_rows,
        col_widths=[5.5, 11.0])

    doc.add_paragraph()
    add_paragraph(doc, "URL triển khai:", bold=True)
    add_bullet(doc, "https://ocr-r2-worker.ocr-r2-worker.workers.dev")

    doc.add_paragraph()

    # ═══════════════════════════════════════════
    # 3. SỐ LIỆU THỰC TẾ TỪ CƠ SỞ DỮ LIỆU
    # ═══════════════════════════════════════════
    add_heading(doc, "3. SỐ LIỆU THỰC TẾ TỪ CƠ SỞ DỮ LIỆU (D1)", level=1)
    add_paragraph(doc, "Dữ liệu cập nhật tại thời điểm 29/09/2026:")

    stats_rows = [
        ("Tổng số phiếu (documents)", "89 phiếu"),
        ("Tổng dòng chi tiết (document_items)", "838 dòng"),
        ("Số mã gỗ riêng biệt (so_xe)", "78 mã"),
        ("Từ điển viết tắt", "54 từ"),
        ("Tổng khối lượng gỗ tròn", "7.319,832 m³"),
        ("Tổng khối lượng thành khí", "2.244,458 m³"),
        ("Tỷ lệ thu hồi bình quân", "~30,66%"),
        ("Dung lượng D1 sử dụng", "~0,74 MB"),
    ]
    add_table_with_header(doc,
        ["Chỉ số", "Giá trị"],
        stats_rows,
        col_widths=[8.0, 6.0],
        header_color="375623")

    doc.add_paragraph()
    note_p = doc.add_paragraph()
    note_r = note_p.add_run(
        "⚠  Tỷ lệ thu hồi ~30,66% phản ánh hiệu suất chế biến gỗ: "
        "cứ 3 m³ gỗ tròn cho ra khoảng 0,92 m³ gỗ thành khí (sản phẩm thương mại).")
    note_r.font.size = Pt(10)
    note_r.font.italic = True
    note_r.font.color.rgb = RGBColor(165, 42, 42)

    doc.add_paragraph()

    # ═══════════════════════════════════════════
    # 4. MÔ HÌNH DỮ LIỆU
    # ═══════════════════════════════════════════
    add_heading(doc, "4. MÔ HÌNH DỮ LIỆU (DATABASE SCHEMA)", level=1)
    add_paragraph(doc,
        "Cơ sở dữ liệu D1 gồm 5 bảng chính, quan hệ như sau:")

    # Bảng documents
    add_heading(doc, "4.1 Bảng documents", level=2)
    doc_cols = [
        ("id", "INTEGER", "PRIMARY KEY AUTOINCREMENT"),
        ("file_name", "TEXT", "Tên file ảnh gốc"),
        ("upload_time", "DATETIME", "Thời gian upload"),
        ("raw_json", "TEXT", "JSON đầy đủ kết quả OCR (header + items)"),
        ("so_xe", "TEXT", "Mã gỗ (trích từ raw_json)"),
        ("ngay_xe", "TEXT", "Ngày xẻ"),
        ("loai_go", "TEXT", "Loại gỗ"),
        ("kich_thuoc_go_tron", "TEXT", "Kích thước gỗ tròn"),
        ("khoi_luong_go_tron", "REAL", "Khối lượng gỗ tròn (m³)"),
        ("quy_cach_go_tron", "TEXT", "Quy cách gỗ tròn"),
        ("tong_khoi_luong", "REAL", "Tổng KL thành khí tính từ items"),
    ]
    add_table_with_header(doc,
        ["Cột", "Kiểu", "Mô tả"],
        doc_cols,
        col_widths=[4.5, 3.0, 9.0])

    doc.add_paragraph()

    # Bảng document_items
    add_heading(doc, "4.2 Bảng document_items", level=2)
    items_cols = [
        ("id", "INTEGER", "PRIMARY KEY AUTOINCREMENT"),
        ("document_id", "INTEGER", "FK → documents.id"),
        ("so_thu_tu", "INTEGER", "Số thứ tự dòng"),
        ("so_cay", "TEXT", "Số cây gỗ"),
        ("chieu_dai", "REAL", "Chiều dài (m)"),
        ("chieu_rong", "REAL", "Chiều rộng (mm)"),
        ("chieu_cao", "REAL", "Chiều cao / dày (mm)"),
        ("vanh", "TEXT", "Vành (loại vành gỗ)"),
        ("khoi_luong", "REAL", "Khối lượng thành khí (m³)"),
        ("ghi_chu", "TEXT", "Ghi chú"),
    ]
    add_table_with_header(doc,
        ["Cột", "Kiểu", "Mô tả"],
        items_cols,
        col_widths=[4.5, 3.0, 9.0])

    doc.add_paragraph()

    # Bảng don_gia_ma_go
    add_heading(doc, "4.3 Bảng don_gia_ma_go", level=2)
    dg_cols = [
        ("ma_go", "TEXT", "PRIMARY KEY – Mã gỗ (trùng với so_xe)"),
        ("don_gia", "REAL", "Đơn giá gỗ tròn (VNĐ/m³)"),
        ("ghi_chu", "TEXT", "Ghi chú"),
        ("ngay_cap_nhat", "DATETIME", "Ngày cập nhật đơn giá"),
    ]
    add_table_with_header(doc,
        ["Cột", "Kiểu", "Mô tả"],
        dg_cols,
        col_widths=[4.5, 3.0, 9.0])

    doc.add_paragraph()

    # Bảng abbreviations
    add_heading(doc, "4.4 Bảng abbreviations", level=2)
    abbr_cols = [
        ("id", "INTEGER", "PRIMARY KEY AUTOINCREMENT"),
        ("abbreviation", "TEXT", "Viết tắt (VD: đoTT)"),
        ("full_text", "TEXT", "Nghĩa đầy đủ"),
    ]
    add_table_with_header(doc,
        ["Cột", "Kiểu", "Mô tả"],
        abbr_cols,
        col_widths=[4.5, 3.0, 9.0])

    doc.add_paragraph()

    # Bảng system_settings
    add_heading(doc, "4.5 Bảng system_settings", level=2)
    ss_cols = [
        ("key", "TEXT", "PRIMARY KEY – Khóa cài đặt"),
        ("value", "TEXT", "Giá trị"),
        ("updated_at", "DATETIME", "Thời gian cập nhật"),
    ]
    add_table_with_header(doc,
        ["Cột", "Kiểu", "Mô tả"],
        ss_cols,
        col_widths=[4.5, 3.0, 9.0])

    doc.add_paragraph()

    # ═══════════════════════════════════════════
    # 5. LUỒNG XỬ LÝ OCR
    # ═══════════════════════════════════════════
    add_heading(doc, "5. LUỒNG XỬ LÝ OCR", level=1)
    steps = [
        ("Bước 1 – Upload ảnh",
         "Người dùng chọn ảnh phiếu nhật ký từ thiết bị, nhấn 'Nhận dạng'. "
         "Frontend gửi multipart/form-data POST lên /api/ocr."),
        ("Bước 2 – Lưu R2",
         "Worker nhận ảnh, upload lên Cloudflare R2 (Object Storage) "
         "với tên file tự động (timestamp + UUID)."),
        ("Bước 3 – Gọi AI Vision",
         "callVision() dispatcher gọi OpenRouter (Gemini 2.5 Flash). "
         "Nếu lỗi, tự động fallback sang Gemini API trực tiếp. "
         "Prompt được thiết kế đặc thù cho phiếu nhật ký xẻ gỗ."),
        ("Bước 4 – Parse & Validate",
         "parseAndValidateOcrResponse() làm sạch JSON trả về: "
         "tách header (mã gỗ, ngày, đơn giá, tổng KL thành khí) "
         "và mảng items (từng dòng gỗ thành khí)."),
        ("Bước 5 – Lưu D1",
         "saveOrUpdateDocument() ghi vào bảng documents và document_items. "
         "raw_json lưu toàn bộ data gốc để không mất thông tin."),
        ("Bước 6 – Hiển thị & Đối chiếu",
         "Frontend render kết quả OCR vào form, tính tổng KL thành khí từ bảng, "
         "so sánh với ô 'Tổng KL thành khí' OCR đọc được. "
         "Cảnh báo màu đỏ/xanh theo kết quả đối chiếu."),
        ("Bước 7 – Xuất Excel Batch",
         "Người dùng chọn nhiều phiếu → xuất Excel 6 sheet: "
         "DS Phiếu, Chi tiết Gỗ, Tổng hợp, Thống kê, Đơn giá, Từ điển viết tắt."),
    ]

    for title, desc in steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        run_title = p.add_run(f"{title}: ")
        run_title.bold = True
        run_title.font.size = Pt(11)
        run_desc = p.add_run(desc)
        run_desc.font.size = Pt(11)

    doc.add_paragraph()

    # ═══════════════════════════════════════════
    # 6. TÍNH NĂNG NỔI BẬT
    # ═══════════════════════════════════════════
    add_heading(doc, "6. TÍNH NĂNG NỔI BẬT", level=1)
    features = [
        ("AI Vision OCR đa mô hình",
         "Hỗ trợ OpenRouter (Gemini 2.5 Flash, Qwen2.5-VL 72B) "
         "và Google Gemini trực tiếp. Tự động fallback khi mô hình chính lỗi."),
        ("Live Reconciliation (Đối chiếu trực tiếp)",
         "Hệ thống tự động tính tổng KL thành khí từ bảng chi tiết và so sánh "
         "với giá trị OCR đọc từ phiếu. Hiển thị trạng thái khớp/lệch realtime."),
        ("Cảnh báo trước khi xuất Excel",
         "Nếu phát hiện số liệu lệch, hệ thống yêu cầu xác nhận trước khi "
         "xuất Excel Batch để tránh sai sót dữ liệu báo cáo."),
        ("Database Đơn giá theo Mã gỗ",
         "Bảng don_gia_ma_go liên kết đơn giá với từng mã gỗ cụ thể, "
         "tự động tra cứu và điền vào form khi OCR xong."),
        ("Model Badge",
         "Hiển thị tên mô hình AI đã xử lý phiếu (VD: 🤖 openrouter/gemini-2.5-flash) "
         "trực tiếp trên giao diện để người dùng biết nguồn gốc kết quả."),
        ("Excel Batch 6 sheet",
         "Xuất Excel theo lô nhiều phiếu, bao gồm: DS Phiếu, Chi tiết Gỗ thành khí, "
         "Tổng hợp theo loại gỗ, Thống kê KPI, Đơn giá, Từ điển viết tắt."),
        ("Từ điển viết tắt",
         "Quản lý 54+ từ viết tắt đặc thù ngành gỗ, hỗ trợ AI và người dùng "
         "đọc đúng nội dung phiếu (VD: đoTT = đoạn thẳng tính, V = viên)."),
        ("Serverless hoàn toàn",
         "Không cần server, database, hay infrastructure riêng. "
         "Toàn bộ hệ thống chạy trên Cloudflare Workers + D1 + R2."),
    ]

    feat_rows = [(f[0], f[1]) for f in features]
    add_table_with_header(doc,
        ["Tính năng", "Mô tả"],
        feat_rows,
        col_widths=[5.0, 11.5],
        header_color="7B2D8B")

    doc.add_paragraph()

    # ═══════════════════════════════════════════
    # 7. CÔNG NGHỆ SỬ DỤNG
    # ═══════════════════════════════════════════
    add_heading(doc, "7. CÔNG NGHỆ SỬ DỤNG", level=1)
    tech_rows = [
        ("Cloudflare Workers", "JavaScript/Node.js", "Runtime Serverless"),
        ("Cloudflare D1", "SQLite", "Cơ sở dữ liệu quan hệ"),
        ("Cloudflare R2", "S3-compatible", "Lưu trữ file ảnh"),
        ("OpenRouter", "REST API", "AI Gateway đa mô hình"),
        ("Gemini 2.5 Flash", "Vision LLM", "Mô hình OCR chính"),
        ("Qwen2.5-VL 72B", "Vision LLM", "Mô hình OCR thay thế"),
        ("ExcelJS", "Node.js Library", "Tạo file Excel"),
        ("Python python-docx", "Python Library", "Tạo file báo cáo DOCX"),
        ("Wrangler CLI", "Cloudflare Tool", "Build và Deploy Worker"),
    ]
    add_table_with_header(doc,
        ["Công nghệ", "Loại", "Mục đích"],
        tech_rows,
        col_widths=[5.0, 4.5, 7.0],
        header_color="C55A11")

    doc.add_paragraph()

    # ═══════════════════════════════════════════
    # 8. CẤU TRÚC FILE DỰ ÁN
    # ═══════════════════════════════════════════
    add_heading(doc, "8. CẤU TRÚC FILE DỰ ÁN", level=1)
    file_rows = [
        ("cloudflare_worker/src/index.js", "Router chính – xử lý tất cả API endpoints"),
        ("cloudflare_worker/src/ocr.js", "AI Vision logic – prompt, gọi API, parse kết quả"),
        ("cloudflare_worker/src/db.js", "Toàn bộ truy vấn D1 – lưu/đọc documents, items"),
        ("cloudflare_worker/src/excel.js", "ExcelJS – tạo 6 sheet báo cáo"),
        ("cloudflare_worker/src/html.js", "Re-export index.html để bundle vào Worker"),
        ("cloudflare_worker/src/index.html", "Sync từ templates/index.html trước mỗi deploy"),
        ("cloudflare_worker/wrangler.toml", "Cấu hình Cloudflare: AI_PROVIDER, model, D1, R2"),
        ("templates/index.html", "File nguồn giao diện SPA – source of truth"),
        ("generate_report.py", "Script tạo báo cáo DOCX này"),
        (".env", "Biến môi trường local (AI_PROVIDER, API keys)"),
    ]
    add_table_with_header(doc,
        ["File", "Mô tả"],
        file_rows,
        col_widths=[7.0, 9.5],
        header_color="1F4E79")

    doc.add_paragraph()

    # ═══════════════════════════════════════════
    # 9. QUY TRÌNH DEPLOY
    # ═══════════════════════════════════════════
    add_heading(doc, "9. QUY TRÌNH DEPLOY", level=1)
    deploy_steps = [
        "Chỉnh sửa giao diện trong templates/index.html",
        "Copy sang cloudflare_worker/src/:",
        "    Copy-Item templates\\index.html cloudflare_worker\\src\\index.html",
        "Deploy lên Cloudflare:",
        "    cd cloudflare_worker",
        "    wrangler deploy",
        "Kiểm tra tại: https://ocr-r2-worker.ocr-r2-worker.workers.dev",
    ]
    for s in deploy_steps:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        r = p.add_run(s)
        if s.strip().startswith("Copy-Item") or s.strip().startswith("wrangler") or s.strip().startswith("cd "):
            r.font.name = "Courier New"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0, 100, 0)
        else:
            r.font.size = Pt(11)

    doc.add_paragraph()

    # ═══════════════════════════════════════════
    # 10. LƯU Ý QUAN TRỌNG
    # ═══════════════════════════════════════════
    add_heading(doc, "10. LƯU Ý QUAN TRỌNG", level=1)
    notes = [
        "Luôn sync templates/index.html → cloudflare_worker/src/index.html trước mỗi lần wrangler deploy.",
        "max_tokens phải set là 4096 trong OpenRouter payload (tránh lỗi 402 credit limit với default 65535).",
        "Trường don_gia và tong_khoi_luong_thanh_khi KHÔNG có cột riêng trong D1, chỉ lưu trong raw_json.",
        "Model slug phải đúng: google/gemini-2.5-flash (không phải gemini-2.0-flash-001).",
        "Tên biến nội bộ (so_xe, kich_thuoc_go_tron, ...) giữ nguyên 100% – chỉ đổi tên ở lớp UI/label.",
    ]
    for n in notes:
        add_bullet(doc, f"⚠  {n}")

    doc.add_paragraph()

    # ═══════════════════════════════════════════
    # FOOTER
    # ═══════════════════════════════════════════
    doc.add_paragraph()
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_r = footer_p.add_run(
        f"— Báo cáo được tạo tự động bởi generate_report.py · {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')} —")
    footer_r.font.size = Pt(9)
    footer_r.font.color.rgb = RGBColor(130, 130, 130)
    footer_r.font.italic = True

    # ── Save ──
    out_path = "bao_cao_tong_quan_ocr_go.docx"
    doc.save(out_path)
    print(f"[OK] Da tao: {out_path}")
    return out_path


if __name__ == "__main__":
    build_report()
